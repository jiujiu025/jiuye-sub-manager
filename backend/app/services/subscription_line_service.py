"""客户线路查询和 Excel 导出服务。"""

from __future__ import annotations

from io import BytesIO
from math import ceil
from urllib.parse import quote

import qrcode
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.core.exceptions import BusinessError
from app.exporters import export_nodes
from app.models.customer import Customer
from app.models.node import Node
from app.models.package import Package
from app.models.subscription import CustomerSubscription
from app.models.three_xui import ForwardedNodeMapping
from app.models.user import User
from app.schemas.subscription_line import (
    SubscriptionLineExportRequest,
    SubscriptionLinePage,
    SubscriptionLineSummary,
)
from app.services.manager_scope import subscription_manager_clause
from app.services.package_service import PackageService
from app.services.subscription_state import is_effectively_active


class SubscriptionLineService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def list_lines(
        self,
        *,
        manager_user_id: int | None,
        page: int = 1,
        page_size: int = 50,
        keyword: str | None = None,
        region: str | None = None,
        status: str | None = None,
    ) -> SubscriptionLinePage:
        rows = self._build_lines(manager_user_id=manager_user_id)
        rows = self._filter(rows, keyword=keyword, region=region, status=status)
        total = len(rows)
        safe_page = max(1, page)
        safe_size = max(1, min(page_size, 200))
        start = (safe_page - 1) * safe_size
        return SubscriptionLinePage(
            items=rows[start : start + safe_size],
            total=total,
            page=safe_page,
            page_size=safe_size,
            pages=max(1, ceil(total / safe_size)) if total else 0,
        )

    def export_excel(
        self,
        payload: SubscriptionLineExportRequest,
        *,
        manager_user_id: int | None,
        is_admin: bool = False,
    ) -> bytes:
        rows = self._build_lines(manager_user_id=manager_user_id)
        rows = self._filter(rows, keyword=payload.keyword, region=payload.region, status=payload.status)
        if not is_admin:
            rows = [row for row in rows if row.status == "active"]
        if payload.line_keys:
            selected = set(payload.line_keys)
            rows = [row for row in rows if row.line_key in selected]
        if not rows:
            raise BusinessError("没有符合条件的可导出线路")

        node_ids = {row.node_id for row in rows}
        export_nodes_by_id = {
            node.id: node
            for node in self.db.scalars(select(Node).where(Node.id.in_(node_ids))).all()
        }
        export_mappings = list(
            self.db.scalars(
                select(ForwardedNodeMapping).where(
                    ForwardedNodeMapping.forwarded_node_id.in_(node_ids),
                    ForwardedNodeMapping.status == "active",
                )
            ).all()
        )
        mapping_by_forwarded_id = {
            mapping.forwarded_node_id: mapping for mapping in export_mappings
        }
        source_ids = {mapping.source_node_id for mapping in export_mappings}
        export_nodes_by_id.update(
            {
                node.id: node
                for node in self.db.scalars(select(Node).where(Node.id.in_(source_ids))).all()
            }
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "线路清单"
        headers = ["时间", "地区", "转发前IP", "转发后IP", "转发后IP的二维码", "客户备注", "到期时间"]
        sheet.append(headers)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill

        for row in rows:
            node = export_nodes_by_id.get(row.node_id)
            forwarded_link = self._line_link(row, node)
            mapping = mapping_by_forwarded_id.get(row.node_id)
            source = None
            if mapping is not None:
                source = export_nodes_by_id.get(mapping.source_node_id)
            original_link = self._line_link(row, source or node)
            sheet.append(
                [
                    row.node_created_at.replace(tzinfo=None),
                    row.customer_region or row.node_country or "",
                    original_link,
                    forwarded_link,
                    "",
                    row.customer_note or "",
                    row.expires_at.replace(tzinfo=None) if row.expires_at else "长期有效",
                ]
            )
            if forwarded_link:
                qr = qrcode.make(forwarded_link)
                image_buffer = BytesIO()
                qr.save(image_buffer, format="PNG")
                image_buffer.seek(0)
                image = ExcelImage(image_buffer)
                image.width = 96
                image.height = 96
                image.anchor = f"E{sheet.max_row}"
                sheet.add_image(image)
                sheet.row_dimensions[sheet.max_row].height = 78

        widths = {"A": 20, "B": 14, "C": 28, "D": 72, "E": 20, "F": 32, "G": 20}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _build_lines(self, *, manager_user_id: int | None) -> list[SubscriptionLineSummary]:
        statement = (
            select(CustomerSubscription)
            .options(
                joinedload(CustomerSubscription.customer).joinedload(Customer.manager),
                joinedload(CustomerSubscription.package).joinedload(Package.rules),
                joinedload(CustomerSubscription.manager),
                joinedload(CustomerSubscription.token),
            )
            .order_by(CustomerSubscription.started_at.desc(), CustomerSubscription.id.desc())
        )
        if manager_user_id is not None:
            statement = statement.join(CustomerSubscription.customer).where(
                subscription_manager_clause(manager_user_id)
            )
        subscriptions = list(self.db.scalars(statement).unique().all())
        package_service = PackageService(self.db)
        all_nodes = list(self.db.scalars(select(Node)).all())
        enabled_nodes = [node for node in all_nodes if node.enabled]
        node_cache = {node.id: node for node in all_nodes}
        package_preview_cache = {
            package.id: package_service.preview_from_nodes(package, enabled_nodes)
            for package in {subscription.package for subscription in subscriptions}
        }
        mappings = list(
            self.db.scalars(
                select(ForwardedNodeMapping).where(ForwardedNodeMapping.status == "active")
            ).all()
        )
        mapping_cache: dict[int, ForwardedNodeMapping] = {
            mapping.forwarded_node_id: mapping for mapping in mappings
        }
        lines: list[SubscriptionLineSummary] = []
        for subscription in subscriptions:
            for preview in package_preview_cache.get(subscription.package.id, []):
                node = node_cache.get(preview.id)
                if node is None:
                    continue
                mapping = mapping_cache.get(node.id)
                source = node_cache.get(mapping.source_node_id) if mapping else None
                manager = subscription.manager or subscription.customer.manager
                lines.append(
                    SubscriptionLineSummary(
                        line_key=f"{subscription.id}:{node.id}",
                        line_code=f"L-{subscription.id}-{node.node_code or node.id}",
                        subscription_id=subscription.id,
                        customer_id=subscription.customer_id,
                        customer_name=subscription.customer.name,
                        customer_region=subscription.customer.region,
                        customer_note=subscription.customer.note or subscription.note,
                        package_name=subscription.package.name,
                        manager_username=manager.username if manager else None,
                        node_id=node.id,
                        node_code=node.node_code,
                        node_name=preview.name,
                        node_type=node.type,
                        node_country=node.country,
                        source_server=source.server if source else None,
                        forwarded=mapping is not None,
                        forwarded_server=node.server,
                        started_at=subscription.started_at,
                        node_created_at=node.created_at,
                        expires_at=subscription.expires_at,
                        status=self._status(subscription),
                    )
                )
        lines.sort(key=lambda item: (-item.node_created_at.timestamp(), item.line_code))
        return lines

    @staticmethod
    def _filter(rows, *, keyword: str | None, region: str | None, status: str | None):
        query = (keyword or "").strip().lower()
        return [
            row
            for row in rows
            if (not query or query in " ".join(
                [row.customer_name, row.package_name, row.node_name, row.line_code]
            ).lower())
            and (not region or row.customer_region == region)
            and (not status or row.status == status)
        ]

    @staticmethod
    def _status(subscription: CustomerSubscription) -> str:
        if subscription.token and subscription.token.revoked_at:
            return "revoked"
        if subscription.status == "cancelled":
            return "cancelled"
        if subscription.status == "paused":
            return "paused"
        return "active" if is_effectively_active(subscription) else "expired"

    def _line_link(self, row: SubscriptionLineSummary, node: Node | None = None) -> str:
        node = node or self.db.get(Node, row.node_id)
        if node is None:
            return ""
        content, count = export_nodes([(node, row.node_name)], "uri")
        if count != 1:
            return ""
        return content.strip()
