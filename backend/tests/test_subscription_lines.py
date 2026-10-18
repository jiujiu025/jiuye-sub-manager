"""全部线路查询和 Excel 导出测试。"""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _headers(client: TestClient) -> dict[str, str]:
    response = client.post(
        "/api/admin/login",
        json={"username": "admin", "password": "TestPass123!"},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_subscription_lines_can_filter_and_export_excel_with_qr(
    client: TestClient,
) -> None:
    headers = _headers(client)
    suffix = uuid4().hex[:8]
    node = client.post(
        "/api/nodes",
        json={
            "name": f"线路节点-{suffix}",
            "type": "vless",
            "server": f"line-{suffix}.example.test",
            "port": 443,
            "uuid": str(uuid4()),
        },
        headers=headers,
    )
    assert node.status_code == 201
    package = client.post(
        "/api/packages",
        json={
            "name": f"线路套餐-{suffix}",
            "rules": {"node_ids": [node.json()["id"]]},
        },
        headers=headers,
    )
    assert package.status_code == 201
    customer = client.post(
        "/api/customers",
        json={"name": f"线路客户-{suffix}", "region": "香港", "note": "线路交付备注"},
        headers=headers,
    )
    assert customer.status_code == 201
    subscription = client.post(
        "/api/subscriptions",
        json={
            "customer_id": customer.json()["id"],
            "package_id": package.json()["id"],
        },
        headers=headers,
    )
    assert subscription.status_code == 201

    listed = client.get(
        "/api/subscription-lines",
        params={"keyword": suffix, "region": "香港"},
        headers=headers,
    )
    assert listed.status_code == 200
    body = listed.json()
    assert body["total"] == 1
    assert body["items"][0]["customer_note"] == "线路交付备注"
    assert "link" not in body["items"][0]

    exported = client.post(
        "/api/subscription-lines/export",
        json={"line_keys": [body["items"][0]["line_key"]]},
        headers=headers,
    )
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    assert [sheet.cell(1, index).value for index in range(1, 8)] == [
        "时间",
        "地区",
        "转发前IP",
        "转发后IP",
        "转发后IP的二维码",
        "客户备注",
        "到期时间",
    ]
    assert sheet.cell(2, 2).value == "香港"
    assert sheet.cell(2, 3).value.startswith("vless://")
    assert sheet.cell(2, 4).value.startswith("vless://")
    assert len(sheet._images) == 1
